from io import BytesIO

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader
from zipfile import BadZipFile, ZipFile
from backend.schemas import MAX_RESUME_TEXT_LENGTH
from backend.services.resume_validation import validate_resume_content

MAX_PDF_PAGES = 50
MAX_DOCX_UNCOMPRESSED_SIZE = 10 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100
DOCX_RATIO_CHECK_MIN_SIZE = 1024 * 1024


def extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def validate_docx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            total_size = 0

            for entry in archive.infolist():
                if entry.is_dir():
                    continue

                total_size += entry.file_size

                if total_size > MAX_DOCX_UNCOMPRESSED_SIZE:
                    raise HTTPException(
                        status_code=413,
                        detail="The DOCX document contains too much uncompressed data.",
                    )

                if (
                    entry.file_size > DOCX_RATIO_CHECK_MIN_SIZE
                    and entry.file_size / max(entry.compress_size, 1)
                    > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise HTTPException(
                        status_code=413,
                        detail="The DOCX document has an unsafe compression ratio.",
                    )

    except HTTPException:
        raise
    except (BadZipFile, OSError, RuntimeError, ValueError) as exc:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is not a valid DOCX document.",
        ) from exc


def extract_resume_text(filename: str, content: bytes) -> str:
    filename = filename.lower()

    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported",
        )

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded resume is empty",
        )

    if filename.endswith(".pdf"):
        if b"%PDF-" not in content[:1024]:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is not a valid PDF",
            )
    elif not content.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is not a valid DOCX document",
        )

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported",
                )

            if len(reader.pages) > MAX_PDF_PAGES:
                raise HTTPException(
                    status_code=413,
                    detail=f"PDF resumes cannot exceed {MAX_PDF_PAGES} pages.",
                )

            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            validate_docx_archive(content)
            text = extract_docx_text(content)

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read the resume file. Check that it is a valid, unprotected PDF or DOCX.",
        ) from exc

    if not text.strip():
        raise HTTPException(
            status_code=400,
            detail="No readable text found in the resume",
        )

    if len(text) > MAX_RESUME_TEXT_LENGTH:
        raise HTTPException(
            status_code=413,
            detail=(
                "The extracted resume exceeds the "
                f"{MAX_RESUME_TEXT_LENGTH:,}-character limit"
            ),
        )
    validate_resume_content(text)
    return text
