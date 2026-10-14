from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab import __file__ as reportlab_package_file
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

SECTION_HEADINGS = {
    "SUMMARY",
    "PROFESSIONAL SUMMARY",
    "PROFILE",
    "SKILLS",
    "TECHNICAL SKILLS",
    "CORE SKILLS",
    "EXPERIENCE",
    "WORK EXPERIENCE",
    "PROFESSIONAL EXPERIENCE",
    "EMPLOYMENT HISTORY",
    "PROJECTS",
    "PERSONAL PROJECTS",
    "EDUCATION",
    "CERTIFICATIONS",
    "CERTIFICATES",
    "ACHIEVEMENTS",
    "AWARDS",
    "PUBLICATIONS",
    "LANGUAGES",
    "VOLUNTEER EXPERIENCE",
    "INTERESTS",
}

FONT_NAME = "Arial"
BODY_FONT_SIZE = 10
HEADING_FONT_SIZE = 11
NAME_FONT_SIZE = 20
CONTACT_FONT_SIZE = 9
PDF_FONT_NAME = "Vera"
PDF_BOLD_FONT_NAME = "Vera-Bold"


def _section_heading(line: str) -> bool:
    return line.strip().rstrip(":").strip().upper() in SECTION_HEADINGS


def _is_bullet(line: str) -> bool:
    stripped = line.lstrip()
    return stripped.startswith(("• ", "- ", "* ", "‣ ", "▪ ", "– "))


def _is_contact_line(line: str) -> bool:
    lowered = line.lower()
    return (
        any(
            marker in lowered
            for marker in (
                "@",
                "linkedin",
                "github",
                "http://",
                "https://",
                "phone:",
                "email:",
            )
        )
        or sum(character.isdigit() for character in line) >= 7
    )


def _register_pdf_fonts() -> None:
    registered_fonts = pdfmetrics.getRegisteredFontNames()
    font_directory = Path(reportlab_package_file).parent / "fonts"

    if PDF_FONT_NAME not in registered_fonts:
        pdfmetrics.registerFont(
            TTFont(PDF_FONT_NAME, str(font_directory / "Vera.ttf"))
        )
    if PDF_BOLD_FONT_NAME not in registered_fonts:
        pdfmetrics.registerFont(
            TTFont(PDF_BOLD_FONT_NAME, str(font_directory / "VeraBd.ttf"))
        )


def generate_docx(resume_text: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_FONT_SIZE)
    normal.font.color.rgb = RGBColor(35, 40, 38)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    first_content_line = True
    for line in resume_text.splitlines():
        paragraph = document.add_paragraph()

        if not line.strip():
            paragraph.paragraph_format.space_after = Pt(4)
            continue

        run = paragraph.add_run(line)
        run.font.name = FONT_NAME

        if first_content_line:
            run.bold = True
            run.font.size = Pt(NAME_FONT_SIZE)
            run.font.color.rgb = RGBColor(37, 77, 59)
            paragraph.paragraph_format.space_after = Pt(2)
            first_content_line = False
        elif _is_contact_line(line):
            run.font.size = Pt(CONTACT_FONT_SIZE)
            run.font.color.rgb = RGBColor(90, 99, 93)
            paragraph.paragraph_format.space_after = Pt(8)
        elif _section_heading(line):
            run.bold = True
            run.font.size = Pt(HEADING_FONT_SIZE)
            run.font.color.rgb = RGBColor(37, 77, 59)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(3)
        elif _is_bullet(line):
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            paragraph.paragraph_format.space_after = Pt(2)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_pdf(resume_text: str) -> bytes:
    _register_pdf_fonts()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName=PDF_FONT_NAME,
        fontSize=BODY_FONT_SIZE,
        leading=13,
        textColor=colors.HexColor("#232826"),
        spaceAfter=3,
    )
    name_style = ParagraphStyle(
        "ResumeName",
        parent=body_style,
        fontName=PDF_BOLD_FONT_NAME,
        fontSize=NAME_FONT_SIZE,
        leading=24,
        textColor=colors.HexColor("#254d3b"),
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=body_style,
        fontSize=CONTACT_FONT_SIZE,
        leading=11,
        textColor=colors.HexColor("#5a635d"),
        spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "ResumeSectionHeading",
        parent=body_style,
        fontName=PDF_BOLD_FONT_NAME,
        fontSize=HEADING_FONT_SIZE,
        leading=14,
        textColor=colors.HexColor("#254d3b"),
        spaceBefore=9,
        spaceAfter=3,
        keepWithNext=True,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=body_style,
        leftIndent=14,
        firstLineIndent=-8,
        spaceAfter=2,
    )
    story = []
    first_content_line = True

    for line in resume_text.splitlines():
        if not line.strip():
            story.append(Spacer(1, 8))
            continue

        if first_content_line:
            style = name_style
            first_content_line = False
        elif _is_contact_line(line):
            style = contact_style
        elif _section_heading(line):
            style = heading_style
        elif _is_bullet(line):
            style = bullet_style
        else:
            style = body_style

        story.append(Paragraph(escape(line), style))

    document.build(story)
    return buffer.getvalue()
